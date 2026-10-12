"""
Smart Condyle — Security Assessment Report Generator
Compiles SAST, DAST, API Inventory, and Dependency scan results into:
  - Vulnerability Test Results/security-review.md
  - Vulnerability Test Results/executive-summary.md
  - Vulnerability Test Results/dependency-report.md
  - Vulnerability Test Results/endpoint-inventory.xlsx
  - Vulnerability Test Results/findings.xlsx (4 required sheets)
"""
import json
import os
import sys
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Windows UTF-8 safety
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8', errors='replace')

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))

def load_json(filename):
    filepath = os.path.join(SCRIPT_DIR, filename)
    if os.path.exists(filepath):
        with open(filepath, 'r', encoding='utf-8') as f:
            return json.load(f)
    return []

def severity_color(severity):
    return {
        'Critical': 'D9534F',  # Red
        'High': 'F0AD4E',      # Orange
        'Medium': 'F7DC6F',    # Gold/Yellow
        'Low': '5CB85C',       # Green
    }.get(severity, 'D6DBDF')

def generate_findings_excel(findings, api_inventory, deps, output_dir):
    wb = Workbook()
    
    header_fill = PatternFill(start_color='1B4F72', end_color='1B4F72', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True, size=11)
    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )

    # ── Sheet 1: Security Findings ──
    ws1 = wb.active
    ws1.title = "Security Findings"
    headers1 = [
        'Finding ID', 'Severity', 'Vulnerability Type', 'File Path', 'Line',
        'Endpoint', 'Description', 'Exploitation Scenario', 'Impact', 'Recommended Fix'
    ]
    ws1.append(headers1)
    for col_idx, cell in enumerate(ws1[1], 1):
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for i, f in enumerate(findings, 1):
        row = [
            f.get('id', f'SC-{i:03d}'),
            f.get('severity', 'Medium'),
            f.get('type', 'Vulnerability'),
            f.get('file', 'smart_condyle_backend/main.py'),
            f.get('line', 'N/A'),
            f.get('endpoint', 'N/A'),
            f.get('description', ''),
            f.get('scenario', 'An attacker manipulates input or directly accesses endpoint.'),
            f.get('impact', ''),
            f.get('recommendation', '')
        ]
        ws1.append(row)
        current_row = i + 1
        sev = f.get('severity', 'Medium')
        sev_cell = ws1.cell(row=current_row, column=2)
        sev_cell.fill = PatternFill(start_color=severity_color(sev), fill_type='solid')
        if sev in ('Critical', 'High'):
            sev_cell.font = Font(color='FFFFFF', bold=True)
        else:
            sev_cell.font = Font(color='000000', bold=True)

        for col_idx in range(1, len(headers1) + 1):
            cell = ws1.cell(row=current_row, column=col_idx)
            cell.border = thin_border
            cell.alignment = Alignment(vertical='top', wrap_text=True)

    ws1.column_dimensions['A'].width = 14
    ws1.column_dimensions['B'].width = 14
    ws1.column_dimensions['C'].width = 30
    ws1.column_dimensions['D'].width = 36
    ws1.column_dimensions['E'].width = 8
    ws1.column_dimensions['F'].width = 24
    ws1.column_dimensions['G'].width = 50
    ws1.column_dimensions['H'].width = 50
    ws1.column_dimensions['I'].width = 45
    ws1.column_dimensions['J'].width = 50

    # ── Sheet 2: Endpoint Inventory ──
    ws2 = wb.create_sheet("Endpoint Inventory")
    headers2 = [
        'Endpoint', 'HTTP Method', 'Authentication Required', 'Expected Roles',
        'Controller / File Path', 'Function Name', 'Request Model', 'Category', 'Risk Level', 'Description'
    ]
    ws2.append(headers2)
    for cell in ws2[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for i, ep in enumerate(api_inventory, 1):
        ws2.append([
            ep.get('endpoint', ''),
            ep.get('method', 'GET'),
            'Yes' if ep.get('auth_required') else 'No',
            ep.get('expected_roles', 'Public'),
            ep.get('controller_path', ep.get('file', '')),
            ep.get('function', ''),
            ep.get('request_model', ''),
            ep.get('category', ''),
            ep.get('risk_level', ''),
            ep.get('description', '')
        ])
        current_row = i + 1
        for col_idx in range(1, len(headers2) + 1):
            cell = ws2.cell(row=current_row, column=col_idx)
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center', wrap_text=True)

    ws2.column_dimensions['A'].width = 25
    ws2.column_dimensions['B'].width = 14
    ws2.column_dimensions['C'].width = 24
    ws2.column_dimensions['D'].width = 35
    ws2.column_dimensions['E'].width = 30
    ws2.column_dimensions['F'].width = 22
    ws2.column_dimensions['G'].width = 20
    ws2.column_dimensions['H'].width = 25
    ws2.column_dimensions['I'].width = 32
    ws2.column_dimensions['J'].width = 45

    # ── Sheet 3: Dependency Vulnerabilities ──
    ws3 = wb.create_sheet("Dependency Vulnerabilities")
    headers3 = ['Package', 'Ecosystem', 'Version', 'Vulnerability Description', 'Severity', 'CVE / CWE', 'Fix Available', 'Remediation']
    ws3.append(headers3)
    for cell in ws3[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for i, d in enumerate(deps, 1):
        ws3.append([
            d.get('package', ''),
            d.get('ecosystem', 'PyPI'),
            d.get('version', ''),
            d.get('vulnerability', ''),
            d.get('severity', 'Low'),
            d.get('cve', 'N/A'),
            d.get('fix_available', 'Yes'),
            d.get('recommendation', '')
        ])
        current_row = i + 1
        sev = d.get('severity', 'Low')
        sev_cell = ws3.cell(row=current_row, column=5)
        sev_cell.fill = PatternFill(start_color=severity_color(sev), fill_type='solid')
        if sev in ('Critical', 'High'):
            sev_cell.font = Font(color='FFFFFF', bold=True)
        for col_idx in range(1, len(headers3) + 1):
            cell = ws3.cell(row=current_row, column=col_idx)
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center', wrap_text=True)

    ws3.column_dimensions['A'].width = 35
    ws3.column_dimensions['B'].width = 18
    ws3.column_dimensions['C'].width = 18
    ws3.column_dimensions['D'].width = 50
    ws3.column_dimensions['E'].width = 14
    ws3.column_dimensions['F'].width = 24
    ws3.column_dimensions['G'].width = 15
    ws3.column_dimensions['H'].width = 50

    # ── Sheet 4: Risk Summary ──
    ws4 = wb.create_sheet("Risk Summary")
    headers4 = ['Severity Category', 'Findings Count', 'Percentage', 'Operational Action Level']
    ws4.append(headers4)
    for cell in ws4[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')

    total = len(findings) or 1
    critical = sum(1 for f in findings if f.get('severity') == 'Critical')
    high = sum(1 for f in findings if f.get('severity') == 'High')
    medium = sum(1 for f in findings if f.get('severity') == 'Medium')
    low = sum(1 for f in findings if f.get('severity') == 'Low')

    ws4.append(['Critical', critical, f'{critical/total*100:.1f}%', 'Block Release — Immediate Patching Required'])
    ws4.append(['High', high, f'{high/total*100:.1f}%', 'Remediate Prior to Production Go-Live'])
    ws4.append(['Medium', medium, f'{medium/total*100:.1f}%', 'Remediate in Current / Next Sprint'])
    ws4.append(['Low', low, f'{low/total*100:.1f}%', 'Technical Debt / Hardening Item'])
    ws4.append(['Total Findings', total, '100%', 'Full SAST & DAST Inventory'])
    ws4.append([])
    
    score = max(10, min(100, 100 - (critical * 8 + high * 3 + medium * 1)))
    ws4.append(['Overall Security Posture Score', f'{score}/100', '', 'Status: Action Required'])

    for row_idx in range(2, 8):
        for col_idx in range(1, 5):
            cell = ws4.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center')

    ws4.column_dimensions['A'].width = 30
    ws4.column_dimensions['B'].width = 18
    ws4.column_dimensions['C'].width = 15
    ws4.column_dimensions['D'].width = 45

    output_file = os.path.join(output_dir, 'findings.xlsx')
    wb.save(output_file)
    print(f"[+] Findings Excel generated successfully: {output_file}")

def generate_endpoint_inventory_excel(api_inventory, output_dir):
    wb = Workbook()
    ws = wb.active
    ws.title = "Endpoint Inventory"
    
    headers = [
        '#', 'Endpoint', 'HTTP Method', 'Authentication Required', 'Expected Roles',
        'Controller / File Path', 'Function Name', 'Request Model', 'Category', 'Risk Level', 'Description'
    ]
    ws.append(headers)
    
    header_fill = PatternFill(start_color='1B4F72', end_color='1B4F72', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True)
    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for i, ep in enumerate(api_inventory, 1):
        ws.append([
            i,
            ep.get('endpoint', ''),
            ep.get('method', 'GET'),
            'Yes' if ep.get('auth_required') else 'No',
            ep.get('expected_roles', 'Public'),
            ep.get('controller_path', ep.get('file', '')),
            ep.get('function', ''),
            ep.get('request_model', ''),
            ep.get('category', ''),
            ep.get('risk_level', ''),
            ep.get('description', '')
        ])
        current_row = i + 1
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center', wrap_text=True)

    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 14
    ws.column_dimensions['D'].width = 24
    ws.column_dimensions['E'].width = 35
    ws.column_dimensions['F'].width = 30
    ws.column_dimensions['G'].width = 22
    ws.column_dimensions['H'].width = 20
    ws.column_dimensions['I'].width = 25
    ws.column_dimensions['J'].width = 32
    ws.column_dimensions['K'].width = 45

    output_file = os.path.join(output_dir, 'endpoint-inventory.xlsx')
    wb.save(output_file)
    print(f"[+] Endpoint Inventory Excel generated: {output_file}")

def generate_security_review_md(findings, dast_findings, output_dir):
    critical = [f for f in findings if f.get('severity') == 'Critical']
    high = [f for f in findings if f.get('severity') == 'High']
    medium = [f for f in findings if f.get('severity') == 'Medium']
    low = [f for f in findings if f.get('severity') == 'Low']
    
    score = max(10, min(100, 100 - (len(critical) * 8 + len(high) * 3 + len(medium) * 1)))

    md = f"""# Comprehensive Security Assessment & Penetration Testing Review

**Target:** Smart Condyle AI Mandibular Condyle Fracture Detection System  
**Backend Framework:** FastAPI (Python 3.11)  
**Database:** MongoDB Atlas / Local MongoDB  
**AI Inference Engine:** TensorFlow Lite CNN (`model/condyle_model.tflite`)  
**Frontend Client:** React Native Expo (Android / iOS / Web)  
**Assessment Date:** {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}  
**Evaluation Standard:** OWASP Top 10 API Security, CWE/SANS Top 25, HIPAA PHI Protection Standards  

---

## 1. Vulnerability Findings Summary

| Severity Level | Finding Count | Status |
| :--- | :--- | :--- |
| **Critical** | {len(critical)} | Immediate Hotfix Required |
| **High** | {len(high)} | Fix Prior to Production Go-Live |
| **Medium** | {len(medium)} | Remediate in Current Development Sprint |
| **Low / Informational** | {len(low)} | Defensive Hardening Recommended |
| **Total Security Issues** | **{len(findings)}** | **Full Audit Inventory** |

**Calculated Security Posture Score:** **`{score}/100`**

---

## 2. Dynamic Application Security Testing (DAST) Verification

The live environment (`http://127.0.0.1:8000`) was tested non-destructively. Dynamic runtime probe outcomes:

| Test Category | Target Endpoint | Severity | Runtime Status | Dynamic Probe Finding |
| :--- | :--- | :--- | :--- | :--- |
"""
    for d in dast_findings:
        md += f"| {d['category']} | `{d['endpoint']}` | **{d['severity']}** | `{d['status']}` | {d['evidence']} |\n"

    md += """
---

## 3. Comprehensive Vulnerability Catalog

"""
    for i, f in enumerate(findings, 1):
        md += f"""### Finding {f.get('id', f'SC-{i:03d}')} — [{f.get('severity', 'Medium')}] {f.get('type', 'Vulnerability')}

- **Severity:** `{f.get('severity', 'Medium')}`
- **Vulnerability Type:** {f.get('type', 'Security Flaw')}
- **File Path:** `{f.get('file', 'N/A')}` (Line {f.get('line', 'N/A')})
- **Endpoint:** `{f.get('endpoint', 'N/A')}`

#### Description
{f.get('description', 'N/A')}

#### Exploitation Scenario
{f.get('scenario', 'An adversary crafts an unauthenticated HTTP request targeting this endpoint, bypassing access controls to manipulate or exfiltrate sensitive data.')}

#### Impact
{f.get('impact', 'High impact on confidentiality, integrity, and regulatory compliance.')}

#### Recommended Fix
{f.get('recommendation', 'Apply validation, authentication middleware, and input sanitization.')}

---

"""

    output_file = os.path.join(output_dir, 'security-review.md')
    with open(output_file, 'w', encoding='utf-8') as f:
        f.write(md)
    print(f"[+] Security Review Markdown saved: {output_file}")

def generate_executive_summary_md(findings, output_dir):
    critical = sum(1 for f in findings if f.get('severity') == 'Critical')
    high = sum(1 for f in findings if f.get('severity') == 'High')
    medium = sum(1 for f in findings if f.get('severity') == 'Medium')
    low = sum(1 for f in findings if f.get('severity') == 'Low')
    score = max(10, min(100, 100 - (critical * 8 + high * 3 + medium * 1)))

    md = f"""# Executive Summary

## Total Findings

Critical: {critical}  
High: {high}  
Medium: {medium}  
Low: {low}  

## Most Critical Risks

1. **Unauthenticated Public Dashboard & Clinical PHI Exposure (`/dashboard`)**: The patient records dashboard is publicly accessible without any credentials or session token, directly exposing patient names, phone numbers, injuries, symptoms, and diagnosis results.
2. **Insecure Password Reset Logic (`/reset-password`)**: The password reset API accepts an email address and overwrites credentials without requiring verification of the 6-digit OTP code sent to the doctor's email, enabling instant account takeover.
3. **Insecure Direct Object Reference (IDOR) & Zero Server-Side Authentication (`/patients/*`, `/history/*`)**: API routes trust client-provided `uid` values in request bodies without validating session or Bearer tokens. Any caller can view, overwrite, or delete records belonging to any doctor.
4. **Plaintext Password Storage & Weak Case-Insensitive Matching (`/register`, `/login`)**: Passwords are saved in plaintext without cryptographic hashing (bcrypt/argon2) and verified using case-insensitive comparisons, violating healthcare data security standards.
5. **Permissive CORS Wildcard (`*`) & Missing Defensive Headers**: The API allows all origins with credentials, exposing the backend to cross-origin abuse and client-side attacks.

## Overall Security Score

{score}/100

---

### Security Engineering Roadmap

| Priority | Timeframe | Action Item |
| :--- | :--- | :--- |
| **P0 (Immediate)** | 24 Hours | Restrict `/dashboard` behind authentication; require OTP token in `/reset-password`. |
| **P0 (Immediate)** | 48 Hours | Implement bcrypt password hashing for all user accounts and disable case-insensitive comparison. |
| **P1 (High)** | Sprint 1 | Implement JWT (JSON Web Token) authentication middleware across all `/patients/*`, `/history/*`, and `/profile/*` endpoints. |
| **P1 (High)** | Sprint 1 | Restrict CORS `allow_origins` to trusted domain names; add HTTP security headers (`nosniff`, `DENY`). |
| **P2 (Medium)** | Sprint 2 | Introduce rate limiting on `/login` and `/send-verification-email`; validate uploaded image MIME types in `/predict`. |
"""

    output_file = os.path.join(output_dir, 'executive-summary.md')
    with open(output_file, 'w', encoding='utf-8') as f:
        f.write(md)
    print(f"[+] Executive Summary saved: {output_file}")

def generate_dependency_report_md(deps, output_dir):
    md = f"""# Dependency Review & Supply Chain Security Report

**Target:** Smart Condyle AI Ecosystem (Python Backend & React Native Frontend)  
**Date:** {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}  

---

## 1. Dependency Vulnerability Inventory

| Package | Ecosystem | Version | Vulnerability & CVE | Severity | Fix Available | Remediation |
| :--- | :--- | :--- | :--- | :--- | :--- | :--- |
"""
    for d in deps:
        md += f"| `{d['package']}` | {d['ecosystem']} | `{d['version']}` | {d['vulnerability']} ({d['cve']}) | **{d['severity']}** | {d['fix_available']} | {d['recommendation']} |\n"

    md += """
---

## 2. Supply-Chain Analysis & Hardening Recommendations

### Python Backend
1. **Pin Exact Dependency Versions:** Specify exact versions in `requirements.txt` (e.g. `fastapi==0.111.0`, `pillow==10.3.0`, `python-multipart==0.0.9`) to prevent automated pull of unvetted major releases.
2. **Automate pip-audit:** Integrate `pip-audit` or GitHub Dependabot in CI/CD pipelines to catch published CVEs before deployments.
3. **Isolate AI Inference:** Run TensorFlow / TFLite in a restricted container sandbox with limited memory and filesystem privileges.

### React Native Frontend
1. **Secure Mobile Credential Storage:** Avoid using `@react-native-async-storage/async-storage` for tokens, medical notes, or user IDs. Use hardware-backed `expo-secure-store` (Android Keystore / iOS Keychain).
2. **Certificate Pinning:** Implement TLS certificate pinning on mobile network requests to protect against adversarial Wi-Fi MitM attacks.
3. **Dependency Auditing:** Execute `npm audit` periodically and update React Native dependencies according to Expo SDK release tracks.
"""

    output_file = os.path.join(output_dir, 'dependency-report.md')
    with open(output_file, 'w', encoding='utf-8') as f:
        f.write(md)
    print(f"[+] Dependency Report saved: {output_file}")

def main():
    findings = load_json('sast-results.json')
    api_inventory = load_json('api-inventory.json')
    dast_findings = load_json('dast-results.json')
    deps = load_json('dependency-results.json')

    output_dir = os.path.join(os.path.dirname(SCRIPT_DIR), 'Vulnerability Test Results')
    os.makedirs(output_dir, exist_ok=True)

    print(f"[*] Generating comprehensive security deliverables in: {output_dir}")
    generate_findings_excel(findings, api_inventory, deps, output_dir)
    generate_endpoint_inventory_excel(api_inventory, output_dir)
    generate_security_review_md(findings, dast_findings, output_dir)
    generate_executive_summary_md(findings, output_dir)
    generate_dependency_report_md(deps, output_dir)

    print("\n" + "="*60)
    print("All Deliverables Successfully Generated")
    print("="*60)
    for fname in os.listdir(output_dir):
        fpath = os.path.join(output_dir, fname)
        size = os.path.getsize(fpath)
        print(f"  - {fname:<30} ({size:,} bytes)")

if __name__ == '__main__':
    main()
